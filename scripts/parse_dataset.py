import openpyxl
import json
import os
import re

supervisors_data = [
    {"emp_id": "GLA107250", "name": "Mr. Narendra Mohan", "phone": "9837356128", "email": "narendra.mohan@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA106248", "name": "Mr. Sachin Sharma", "phone": "8077621113", "email": "sachin.sharma@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA108254", "name": "Dr. Anuj Mangal", "phone": "9897534383", "email": "anuj.mangal@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA115127", "name": "Dr. Dhirendra Prasad Yadav", "phone": "9568207247", "email": "dhirendra.yadav@gla.ac.in", "designation": "Assistant Professor", "aliases": ["Dr. D.P. Yadav", "Dr. Dhirendra Prasad Yadav"]},
    {"emp_id": "GLA114099", "name": "Dr. Mayank Agrawal", "phone": "9897626693", "email": "mayank.agrawal@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA119314", "name": "Dr. Anuj Kumar", "phone": "9997189728", "email": "anuj.kumar@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA119315", "name": "Dr. Vinod Jain", "phone": "9813078438", "email": "vinod.jain@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA119316", "name": "Mr. Navin Kumar Agrawal", "phone": "9411253442", "email": "navin.agrawal@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA123254", "name": "Mr. Sanjeev Agrawal", "phone": "8923015116", "email": "sanjiv.agrawal@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA125203", "name": "Mr. Kriti Bansal", "phone": "8273996456", "email": "kriti.bansal@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA123263", "name": "Mr. Santosh Kumar Swarnkar", "phone": "7310663121", "email": "santosh.swarnkar@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA123288", "name": "Ms. Cheshta Bharadwaj", "phone": "8126796015", "email": "cheshtaa.bhardwaj@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA123290", "name": "Dr. Anil Kumar", "phone": "9897667712", "email": "anil.chanchal@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA123053", "name": "Mr. Jayati Krishna Goswami", "phone": "8171806228", "email": "jayati.goswami@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA124174", "name": "Mr. Satish Kumar Maurya", "phone": "9196116503", "email": "satish.maurya@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA124175", "name": "Mr. Aman Deep Singh", "phone": "7500712613", "email": "Aman.deep@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA124196", "name": "Mr. Koushik Choudhury", "phone": "8910496595", "email": "koushik.choudhury@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA123063", "name": "Dr. Puneet Sharma", "phone": "9897265401", "email": "puneet.sharma@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA125219", "name": "Mr Shivam kumar maurya", "phone": "9336393284", "email": "shivamk.maurya@gla.ac.in", "designation": "Teaching Associate"},
    {"emp_id": "GLA124204", "name": "Mr Ashish kumar", "phone": "9599939837", "email": "kumar.ashish@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA124207", "name": "Ms ROSHNI PATOA", "phone": "8077610485", "email": "roshni.patoa@gla.ac.in", "designation": "Teaching Associate"},
    {"emp_id": "GLA125222", "name": "Dr Arun Singh yadav", "phone": "9305591293", "email": "arunsingh.yadav@gla.ac.in", "designation": "Assistant Professor"},
    {"emp_id": "GLA121233", "name": "Mr.Shivank chauhan", "phone": "9694159338", "email": "shivank.chauhan@gla.ac.in", "designation": "Assistant Professor"},
]

def normalize_name(n):
    return re.sub(r'[^a-zA-Z]', '', n.lower())

def find_supervisor(raw_name):
    if not raw_name:
        return None
    raw_norm = normalize_name(raw_name)
    for s in supervisors_data:
        if raw_norm == normalize_name(s['name']):
            return s
        for alias in s.get('aliases', []):
            if raw_norm == normalize_name(alias):
                return s
    # fuzzy contains
    for s in supervisors_data:
        s_norm = normalize_name(s['name'])
        if s_norm in raw_norm or raw_norm in s_norm:
            return s
    return None

wb = openpyxl.load_workbook('BCA and BCA-DS 3rd Year Research Project Group with Guide Odd Sem 2026-27.xlsx')

teams = []
students = []

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    program = 'BCA - DS' if 'DS' in sheetname.upper() else 'BCA'
    curr_team_code = None
    curr_guide_name = None
    curr_team_number = None

    for r in range(2, ws.max_row + 1):
        g_val = ws.cell(r, 9).value
        guide_val = ws.cell(r, 10).value
        s_no = ws.cell(r, 1).value
        course = ws.cell(r, 2).value or program
        section = ws.cell(r, 3).value
        roll = ws.cell(r, 4).value
        name = ws.cell(r, 5).value
        mob = ws.cell(r, 6).value
        email = ws.cell(r, 7).value
        cpi = ws.cell(r, 8).value

        if g_val is not None:
            g_str = str(g_val).strip()
            # extract number if present
            num_match = re.search(r'\d+', g_str)
            curr_team_number = int(num_match.group(0)) if num_match else len(teams) + 1
            if program == 'BCA - DS':
                curr_team_code = f"DS-{g_str.replace('DS', '').strip()}"
            else:
                curr_team_code = f"BCA-{g_str}"
            
            curr_guide_name = guide_val
            sup = find_supervisor(curr_guide_name)
            
            teams.append({
                "team_code": curr_team_code,
                "team_name": f"Team {curr_team_code}",
                "team_number": curr_team_number,
                "program": program,
                "guide_name": curr_guide_name,
                "supervisor_email": sup['email'] if sup else None,
                "supervisor_emp_id": sup['emp_id'] if sup else None
            })

        if name and roll:
            students.append({
                "roll_no": str(roll).strip(),
                "full_name": str(name).strip(),
                "email": str(email).strip().lower() if email else f"{str(roll).strip()}@gla.ac.in",
                "mobile": str(mob).strip() if mob else "9876543210",
                "cpi": float(cpi) if cpi is not None and str(cpi).strip() != '' else None,
                "course": str(course).strip(),
                "section": str(section).strip() if section else "A",
                "team_code": curr_team_code
            })

os.makedirs('data', exist_ok=True)
payload = {
    "supervisors": supervisors_data,
    "teams": teams,
    "students": students,
    "admin": {
        "email": "admin@codeshastra.edu",
        "full_name": "Dr. Project Incharge (Head Administrator)",
        "phone": "9999988888",
        "role": "admin"
    }
}

with open('data/initial_seed.json', 'w', encoding='utf-8') as f:
    json.dump(payload, f, indent=2)

print(f"Dataset compiled successfully!")
print(f"Supervisors: {len(supervisors_data)}")
print(f"Teams: {len(teams)}")
print(f"Students: {len(students)}")
